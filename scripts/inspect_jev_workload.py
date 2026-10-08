"""Compare request volume and the current byte-based pacing budget, without API calls."""
import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from insights.jev_workflow import evaluate_row, validate_workflow


def inspect(config, rows):
    validate_workflow(config['questions'], config.get('derived'))
    execution = config.get('execution', {})
    seen, sizes, counts = set(), [], []
    reused = 0
    staged = any('dependsOn' in q or 'sourceExcerpt' in q for q in config['questions'])
    for row in rows:
        state = {k: row[k] for k in config['sourceColumns'] if row.get(k) is not None and str(row[k]).strip()}
        key = json.dumps(state, ensure_ascii=False)
        if execution.get('reuseIdenticalRows') and key in seen:
            reused += 1
            continue
        seen.add(key)
        def size_call(context, questions):
            payload = {'model': config.get('model', 'jev-1.13.0'), 'state': context, 'questions': questions}
            sizes.append(len(json.dumps(payload, ensure_ascii=False).encode('utf-8')))
            counts.append(sum(len(q.get('criteria', {})) for q in questions.values()))
            answers = {}
            for name, q in questions.items():
                kind = q['type']
                if kind == 'noul':
                    answers[name] = {'type': 'noul', 'noul': 1}
                    continue
                keys = list(q['criteria']) if kind == 'choice' else [str(i) for i in range(len(q['criteria']))]
                answers[name] = {'type': kind, 'confidence': 1, 'probabilities': {k: int(i == 0) for i, k in enumerate(keys)},
                                 **({'choice': keys[0]} if kind == 'choice' else {'score': 0})}
            return {'answers': answers}
        # Placeholder selections size later-stage contexts. They are not model
        # predictions and must never be presented as classified results.
        evaluate_row(state, config['questions'], config.get('derived', []), size_call)
    return {'rows': len(rows), 'requests': len(sizes), 'reusedRows': reused,
            'decisionMode': 'first-option placeholders for staged sizing; no inference' if staged else 'payload sizing; no inference',
            'requestBytes': sum(sizes), 'meanRequestBytes': round(sum(sizes) / len(sizes)) if sizes else 0,
            'meanOptions': round(sum(counts) / len(counts), 2) if counts else 0,
            'pacingBudgetSeconds': round(sum(max(60 / execution.get('requestsPerMinute', 600),
                                                size / execution.get('tokensPerSecond', 100000)) for size in sizes), 2)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--sample', type=Path, default=ROOT / 'sample_jev/jev_output_sample.xlsx')
    parser.add_argument('configs', nargs='*', type=Path, default=[
        ROOT / 'sample_jev/ebola-feedback-coding.config_v19 (1).json',
        ROOT / 'sample_jev/ebola-feedback-coding.fast.config.json'])
    args = parser.parse_args()
    book = load_workbook(args.sample, read_only=True, data_only=True)
    try:
        cells = iter(book.active.values)
        headers = next(cells)
        rows = [dict(zip(headers, row)) for row in cells]
    finally:
        book.close()
    for path in args.configs:
        print(path.name)
        print(json.dumps(inspect(json.loads(path.read_text(encoding='utf-8')), rows), indent=2))


if __name__ == '__main__':
    main()
